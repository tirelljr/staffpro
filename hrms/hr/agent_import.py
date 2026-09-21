# Copyright (c) 2026, Staff Pro BPO and Contributors
# License: GNU General Public License v3. See license.txt

"""Bulk-import agents from a CSV or Excel sheet, with preview before insert."""

from __future__ import annotations

import base64
import csv
import io
import re
from datetime import date, datetime

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate
from frappe.utils.csvutils import to_csv

from hrms.hr.bpo_employee_labels import BELIZE_EMPLOYEE_BANKS, STAFF_PRO_SALARY_CURRENCY

TEXT_FIELDS = {
	"employee_number",
	"username",
	"cell_number",
	"bank_ac_no",
	"pan_number",
	"social_security_number",
	"emergency_phone_number",
}

COLUMNS = (
	{"fieldname": "first_name", "label": "First Name", "required": True, "example": "Maria", "aliases": ("first", "given_name")},
	{"fieldname": "last_name", "label": "Last Name", "required": True, "example": "Santos", "aliases": ("last", "surname", "family_name")},
	{"fieldname": "gender", "label": "Gender", "required": True, "example": "Female", "aliases": ("sex",), "kind": "link", "options": "Gender"},
	{"fieldname": "date_of_birth", "label": "Date of Birth", "required": True, "example": "1995-04-12", "aliases": ("dob", "birth_date"), "kind": "date"},
	{"fieldname": "date_of_joining", "label": "Date of Joining", "required": True, "example": "2024-03-01", "aliases": ("doj", "hire_date", "joining_date"), "kind": "date"},
	{"fieldname": "company", "label": "Company", "required": False, "example": "Staff Pro BPO", "aliases": (), "kind": "link", "options": "Company"},
	{"fieldname": "employee_number", "label": "Employee Number", "required": False, "example": "EMP-00021", "aliases": ("employee_id", "agent_id", "id")},
	{"fieldname": "username", "label": "Username", "required": False, "example": "MSantos", "aliases": ("login", "user_id", "user_name")},
	{"fieldname": "company_email", "label": "Company Email", "required": False, "example": "maria.santos@staffpro.com", "aliases": ("email", "work_email")},
	{"fieldname": "personal_email", "label": "Personal Email", "required": False, "example": "maria.santos@gmail.com", "aliases": ("home_email",)},
	{"fieldname": "cell_number", "label": "Cell Number", "required": False, "example": "+501 610-1234", "aliases": ("mobile", "phone", "cell", "mobile_no")},
	{"fieldname": "status", "label": "Status", "required": False, "example": "Active", "aliases": ()},
	{"fieldname": "department", "label": "Team", "required": False, "example": "Operations", "aliases": ("team", "dept"), "kind": "link", "options": "Department"},
	{"fieldname": "designation", "label": "Role", "required": False, "example": "Agent", "aliases": ("role", "job_title", "title"), "kind": "link", "options": "Designation"},
	{"fieldname": "grade", "label": "Campaign", "required": False, "example": "Collections", "aliases": ("campaign", "employee_grade"), "kind": "link", "options": "Employee Grade"},
	{"fieldname": "branch", "label": "Branch", "required": False, "example": "Belize City", "aliases": (), "kind": "link", "options": "Branch"},
	{"fieldname": "employment_type", "label": "Employment Type", "required": False, "example": "Full-time", "aliases": (), "kind": "link", "options": "Employment Type"},
	{"fieldname": "default_shift", "label": "Default Shift", "required": False, "example": "Day Shift", "aliases": ("shift", "shift_type"), "kind": "link", "options": "Shift Type"},
	{"fieldname": "holiday_list", "label": "Holiday List", "required": False, "example": "Staff Pro Holidays", "aliases": (), "kind": "link", "options": "Holiday List"},
	{"fieldname": "reports_to", "label": "Reports To", "required": False, "example": "James Rivera", "aliases": ("manager", "supervisor"), "kind": "employee"},
	{"fieldname": "ctc", "label": "Agent Hourly", "required": False, "example": "8.50", "aliases": ("hourly", "hourly_rate", "agent_hourly", "pay_rate"), "kind": "float"},
	{"fieldname": "overtime_threshold_hours", "label": "OT Threshold (Hours)", "required": False, "example": "80", "aliases": ("ot_threshold", "ot_threshold_hours", "overtime_threshold"), "kind": "float"},
	{"fieldname": "salary_mode", "label": "Salary Mode", "required": False, "example": "Bank", "aliases": ("pay_mode",)},
	{"fieldname": "bank_name", "label": "Bank Name", "required": False, "example": "Belize Bank", "aliases": ("bank",)},
	{"fieldname": "bank_ac_no", "label": "Bank Account No", "required": False, "example": "1234567890", "aliases": ("bank_account_no", "account_no", "account_number", "bank_account")},
	{"fieldname": "bank_account_type", "label": "Bank Account Type", "required": False, "example": "Checking", "aliases": ("account_type",)},
	{"fieldname": "pan_number", "label": "Tax Number", "required": False, "example": "123456789", "aliases": ("tax_number", "tax_id", "tin")},
	{"fieldname": "social_security_number", "label": "Social Security Number", "required": False, "example": "001234567", "aliases": ("ss_number", "ssn", "ssb_number")},
	{"fieldname": "bill_to_customer", "label": "Bill To Client", "required": False, "example": "Acme Corp", "aliases": ("client", "customer", "bill_to_client"), "kind": "link", "options": "Customer"},
	{"fieldname": "billing_rate", "label": "Billing Rate (Hourly)", "required": False, "example": "18.00", "aliases": ("client_rate", "billing_rate_hourly"), "kind": "float"},
	{"fieldname": "person_to_be_contacted", "label": "Emergency Contact Name", "required": False, "example": "Luis Santos", "aliases": ("emergency_contact", "emergency_contact_name")},
	{"fieldname": "emergency_phone_number", "label": "Emergency Phone", "required": False, "example": "+501 610-5678", "aliases": ("emergency_phone", "emergency_number")},
	{"fieldname": "current_address", "label": "Current Address", "required": False, "example": "12 Albert Street, Belize City", "aliases": ("address", "home_address")},
)

GENDER_ALIASES = {
	"m": "Male",
	"male": "Male",
	"f": "Female",
	"female": "Female",
	"o": "Other",
	"other": "Other",
	"nonbinary": "Other",
	"non_binary": "Other",
}
PREVIEW_FIELDS = (
	"first_name",
	"last_name",
	"company_email",
	"department",
	"designation",
	"grade",
	"ctc",
	"status",
)


def normalize_header(value: str | None) -> str:
	text = re.sub(r"[^a-z0-9]+", "_", (value or "").strip().lower()).strip("_")
	return text


def _column_alias_map() -> dict[str, str]:
	mapping = {}
	for column in COLUMNS:
		fieldname = column["fieldname"]
		keys = {normalize_header(fieldname), normalize_header(column["label"]), *(column.get("aliases") or ())}
		for key in keys:
			if key:
				mapping[key] = fieldname
	return mapping


COLUMN_ALIASES = _column_alias_map()
BANK_ALIASES = {normalize_header(name): name for name in BELIZE_EMPLOYEE_BANKS}
BANK_ALIASES.update(
	{
		"heritage": "Heritage Bank",
		"belize": "Belize Bank",
		"atlantic": "Atlantic Bank",
		"nbb": "National Bank of Belize",
		"national": "National Bank of Belize",
	}
)


def _assert_can_import():
	if not frappe.has_permission("Employee", "create"):
		frappe.throw(_("Not permitted to import agents."), frappe.PermissionError)


@frappe.whitelist()
def get_import_schema() -> dict:
	_assert_can_import()
	naming = frappe.db.get_single_value("HR Settings", "emp_created_by") or "Full Name"
	columns = []
	for column in COLUMNS:
		item = dict(column)
		item["required"] = bool(column["required"] or (column["fieldname"] == "employee_number" and naming == "Employee Number"))
		columns.append(item)
	return {
		"columns": columns,
		"naming": naming,
		"banks": list(BELIZE_EMPLOYEE_BANKS),
		"example": {column["fieldname"]: column["example"] for column in COLUMNS},
	}


@frappe.whitelist()
def download_template(file_format: str = "xlsx"):
	_assert_can_import()
	schema = get_import_schema()
	headers = [column["label"] for column in schema["columns"]]
	example = [column["example"] for column in schema["columns"]]
	fmt = (file_format or "xlsx").strip().lower()
	if fmt == "csv":
		frappe.response["filename"] = "Agent_Import_Template.csv"
		frappe.response["filecontent"] = to_csv([headers, example])
		frappe.response["type"] = "binary"
		return

	from io import BytesIO

	import xlsxwriter

	buffer = BytesIO()
	workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
	sheet = workbook.add_worksheet("Agents")
	header_fmt = workbook.add_format({"bold": True, "bg_color": "#0F1B2D", "font_color": "#FFFFFF", "text_wrap": True})
	text_fmt = workbook.add_format({"num_format": "@"})
	example_fmt = workbook.add_format({"italic": True, "font_color": "#6B7280", "num_format": "@"})
	sheet.freeze_panes(1, 0)
	sheet.autofilter(0, 0, 0, len(headers) - 1)
	for idx, (column, header, sample) in enumerate(zip(schema["columns"], headers, example)):
		width = max(14, min(28, len(header) + 4))
		is_text = column["fieldname"] in TEXT_FIELDS or column.get("kind") in {None, "date", "link", "employee"}
		sheet.set_column(idx, idx, width, text_fmt if is_text else None)
		sheet.write(0, idx, header, header_fmt)
		sheet.write(1, idx, sample, example_fmt)
	workbook.close()
	buffer.seek(0)
	frappe.response["filename"] = "Agent_Import_Template.xlsx"
	frappe.response["filecontent"] = buffer.getvalue()
	frappe.response["type"] = "binary"


@frappe.whitelist()
def preview_agent_import(filename: str | None = None, filedata: str | None = None, file_url: str | None = None) -> dict:
	_assert_can_import()
	headers, body = read_import_table(filename, filedata, file_url)
	return build_preview(headers, body)


@frappe.whitelist()
def import_agents(
	filename: str | None = None,
	filedata: str | None = None,
	file_url: str | None = None,
	skip_errors: int | str | None = 1,
) -> dict:
	_assert_can_import()
	preview = preview_agent_import(filename=filename, filedata=filedata, file_url=file_url)
	ready = [row for row in preview["rows"] if row.get("ok")]
	blocked = [row for row in preview["rows"] if not row.get("ok")]
	if not ready:
		frappe.throw(_("No valid agent rows to import. Fix the errors shown in the preview."))
	if blocked and not cint(skip_errors):
		frappe.throw(_("Fix rows with errors before importing, or skip those rows."))

	created = []
	failed = []
	pending_reports = []
	for row in ready:
		try:
			name = _insert_agent(row["values"], pending_reports)
			created.append({"row_number": row["row_number"], "name": name, "employee_name": row["values"].get("employee_name")})
		except Exception:
			frappe.log_error(title="Agent import failed", message=frappe.get_traceback())
			failed.append({"row_number": row["row_number"], "errors": [_("Could not create this agent. Check the error log.")]})

	_apply_reports_to(pending_reports)
	return {
		"created": created,
		"failed": failed,
		"skipped": [{"row_number": row["row_number"], "errors": row.get("errors") or []} for row in blocked],
		"created_count": len(created),
		"failed_count": len(failed),
		"skipped_count": len(blocked),
	}


def read_import_table(filename: str | None = None, filedata: str | None = None, file_url: str | None = None) -> tuple[list[str], list[list]]:
	name, content = _load_bytes(filename, filedata, file_url)
	ext = (name.rsplit(".", 1)[-1] if "." in name else "").lower()
	if ext in {"xlsx", "xlsm"}:
		rows = _read_xlsx(content)
	elif ext in {"csv", "txt"} or not ext:
		rows = _read_csv(content)
	else:
		frappe.throw(_("Upload a .csv or .xlsx file."))
	if not rows:
		frappe.throw(_("The file is empty."))
	headers = [_cell_str(value) for value in rows[0]]
	if not any(headers):
		frappe.throw(_("The first row must contain column headers."))
	body = []
	for row in rows[1:]:
		values = [_cell_str(value) for value in row]
		if len(values) < len(headers):
			values.extend([""] * (len(headers) - len(values)))
		if not any(values[: len(headers)]):
			continue
		body.append(values[: len(headers)])
	if not body:
		frappe.throw(_("Add at least one agent row under the header row."))
	return headers, body


def map_headers(headers: list[str]) -> tuple[dict[int, str], list[str]]:
	mapped = {}
	seen = set()
	unknown = []
	for idx, header in enumerate(headers):
		if not header:
			continue
		fieldname = COLUMN_ALIASES.get(normalize_header(header))
		if not fieldname:
			unknown.append(header)
			continue
		if fieldname in seen:
			continue
		mapped[idx] = fieldname
		seen.add(fieldname)
	return mapped, unknown


def build_preview(headers: list[str], body: list[list]) -> dict:
	mapping, unknown = map_headers(headers)
	matched = {fieldname: headers[idx] for idx, fieldname in mapping.items()}
	schema = get_import_schema()
	missing_headers = [
		column["label"] for column in schema["columns"] if column["required"] and column["fieldname"] not in matched
	]
	file_names = {_full_name(row, mapping) for row in body}
	preview_rows = []
	seen_names = {}
	seen_emails = {}
	seen_usernames = {}
	for offset, raw in enumerate(body, start=2):
		values = {fieldname: raw[idx].strip() if idx < len(raw) else "" for idx, fieldname in mapping.items()}
		normalized, row_errors = normalize_agent_row(values, file_names=file_names)
		errors = list(row_errors)
		full_name = (normalized.get("employee_name") or "").strip().lower()
		if full_name:
			if full_name in seen_names:
				errors.append(_("This name is used again in row {0}.").format(seen_names[full_name]))
			else:
				seen_names[full_name] = offset
		for fieldname, label in (("company_email", _("Company Email")), ("personal_email", _("Personal Email"))):
			email = (normalized.get(fieldname) or "").strip().lower()
			if not email:
				continue
			if email in seen_emails:
				errors.append(_("{0} is used again in row {1}.").format(label, seen_emails[email]))
			else:
				seen_emails[email] = offset
		username = (normalized.get("username") or "").strip().lower()
		if username:
			if username in seen_usernames:
				errors.append(_("Username is used again in row {0}.").format(seen_usernames[username]))
			else:
				seen_usernames[username] = offset
		if missing_headers:
			errors.insert(0, _("Match the required column headers before importing."))
		preview_rows.append(
			{
				"row_number": offset,
				"ok": not errors,
				"errors": errors,
				"values": normalized,
				"preview": {key: normalized.get(key) or "" for key in PREVIEW_FIELDS},
			}
		)

	ready = sum(1 for row in preview_rows if row["ok"])
	return {
		"headers": headers,
		"matched": matched,
		"unknown_headers": unknown,
		"missing_headers": missing_headers,
		"columns": schema["columns"],
		"rows": preview_rows,
		"ready_count": ready,
		"error_count": len(preview_rows) - ready,
		"total_count": len(preview_rows),
		"naming": schema["naming"],
	}


def normalize_agent_row(values: dict, file_names: set[str] | None = None) -> tuple[dict, list[str]]:
	errors = []
	out = dict(values)
	meta = frappe.get_meta("Employee")
	company = _resolve_company(out.get("company"))
	if out.get("company") and not company:
		errors.append(_("Company {0} was not found.").format(frappe.bold(out.get("company"))))
	out["company"] = company or _default_company()
	if not out["company"]:
		errors.append(_("Set a Company, or choose a default company for your user."))

	for column in COLUMNS:
		fieldname = column["fieldname"]
		if fieldname not in out:
			out[fieldname] = ""
		if not meta.has_field(fieldname) and fieldname not in {"username"}:
			if out.get(fieldname):
				out[fieldname] = ""
			continue
		kind = column.get("kind")
		raw = (out.get(fieldname) or "").strip()
		if column["required"] and not raw:
			if fieldname in values:
				errors.append(_("{0} is required.").format(column["label"]))
			continue
		if not raw:
			continue
		if kind == "date":
			try:
				out[fieldname] = getdate(raw).isoformat()
			except Exception:
				errors.append(_("{0} must be a date (YYYY-MM-DD).").format(column["label"]))
		elif kind == "float":
			out[fieldname] = flt(raw)
		elif fieldname == "gender":
			out[fieldname] = _resolve_gender(raw)
			if not out[fieldname]:
				errors.append(_("Gender {0} was not found. Use Male, Female, or Other.").format(frappe.bold(raw)))
		elif fieldname == "bank_name":
			out[fieldname] = BANK_ALIASES.get(normalize_header(raw), raw)
			if out[fieldname] not in BELIZE_EMPLOYEE_BANKS:
				errors.append(
					_("Bank Name must be one of: {0}.").format(", ".join(BELIZE_EMPLOYEE_BANKS))
				)
		elif fieldname == "bank_account_type":
			choice = raw.title() if raw.lower() in {"checking", "savings"} else raw
			if choice not in {"Checking", "Savings"}:
				errors.append(_("Bank Account Type must be Checking or Savings."))
			out[fieldname] = choice
		elif fieldname == "status":
			choice = raw.title()
			if choice not in {"Active", "Inactive", "Suspended", "Left"}:
				errors.append(_("Status must be Active, Inactive, Suspended, or Left."))
			out[fieldname] = choice
		elif fieldname == "salary_mode":
			choice = raw.title() if raw.lower() != "bank" else "Bank"
			if choice not in {"Bank", "Cash", "Cheque"}:
				errors.append(_("Salary Mode must be Bank, Cash, or Cheque."))
			out[fieldname] = choice
		elif kind == "link":
			resolved = _resolve_link(column["options"], raw, company=out.get("company"))
			if not resolved:
				errors.append(_("{0} {1} was not found.").format(column["label"], frappe.bold(raw)))
			out[fieldname] = resolved or raw
		elif kind == "employee":
			resolved = _resolve_employee(raw, company=out.get("company"))
			if resolved:
				out[fieldname] = resolved
			elif file_names and raw.strip().lower() in {name.lower() for name in file_names if name}:
				out["_reports_to_name"] = raw.strip()
			else:
				errors.append(_("Reports To {0} was not found.").format(frappe.bold(raw)))

	out["first_name"] = (out.get("first_name") or "").strip()
	out["last_name"] = (out.get("last_name") or "").strip()
	out["employee_name"] = _full_name_from_values(out)
	out["status"] = out.get("status") or "Active"
	if out.get("company_email"):
		out["prefered_contact_email"] = "Company Email"
		out["prefered_email"] = out["company_email"]
	elif out.get("personal_email"):
		out["prefered_contact_email"] = "Personal Email"
		out["prefered_email"] = out["personal_email"]
	if out.get("bank_name") or out.get("bank_ac_no"):
		out["salary_mode"] = out.get("salary_mode") or "Bank"
	if meta.has_field("salary_currency"):
		out["salary_currency"] = STAFF_PRO_SALARY_CURRENCY
	if out.get("username"):
		out["user_id"] = out["username"]

	_check_uniques(out, errors)
	return out, errors


def _insert_agent(values: dict, pending_reports: list) -> str:
	payload = {"doctype": "Employee"}
	meta = frappe.get_meta("Employee")
	skip = {"username", "_reports_to_name"}
	for fieldname, value in values.items():
		if fieldname in skip or value in (None, ""):
			continue
		if fieldname == "reports_to" and values.get("_reports_to_name"):
			continue
		if meta.has_field(fieldname):
			payload[fieldname] = value
	if values.get("employee_number") and meta.has_field("naming_series"):
		payload.setdefault("naming_series", "HR-EMP-")
	doc = frappe.get_doc(payload)
	doc.insert()
	if values.get("_reports_to_name"):
		pending_reports.append({"name": doc.name, "reports_to_name": values["_reports_to_name"], "company": doc.company})
	return doc.name


def _apply_reports_to(pending_reports: list[dict]):
	for item in pending_reports:
		manager = _resolve_employee(item["reports_to_name"], company=item.get("company"))
		if not manager:
			continue
		frappe.db.set_value("Employee", item["name"], "reports_to", manager, update_modified=False)


def _check_uniques(values: dict, errors: list[str]):
	name = values.get("employee_name")
	if name and frappe.db.exists("Employee", {"employee_name": name}):
		errors.append(_("An agent named {0} already exists.").format(frappe.bold(name)))
	number = values.get("employee_number")
	if number and frappe.db.exists("Employee", {"employee_number": number}):
		errors.append(_("Employee Number {0} is already used.").format(frappe.bold(number)))
	for fieldname, label in (("company_email", _("Company Email")), ("personal_email", _("Personal Email"))):
		email = (values.get(fieldname) or "").strip()
		if email and frappe.db.exists("Employee", {fieldname: email}):
			errors.append(_("{0} {1} is already used.").format(label, frappe.bold(email)))
	username = (values.get("username") or "").strip()
	if username:
		taken_user = frappe.db.exists("User", {"username": username})
		taken_name = frappe.db.exists("User", username)
		if taken_user or taken_name:
			errors.append(_("Username {0} is already taken.").format(frappe.bold(username)))


def _resolve_company(value: str | None) -> str | None:
	raw = (value or "").strip()
	if not raw:
		return None
	if frappe.db.exists("Company", raw):
		return raw
	return frappe.db.get_value("Company", {"company_name": raw}, "name")


def _default_company() -> str | None:
	return frappe.defaults.get_user_default("Company") or frappe.db.get_single_value("Global Defaults", "default_company")


def _resolve_gender(value: str) -> str | None:
	raw = (value or "").strip()
	if not raw:
		return None
	mapped = GENDER_ALIASES.get(normalize_header(raw), raw)
	if frappe.db.exists("Gender", mapped):
		return mapped
	if frappe.db.exists("Gender", raw):
		return raw
	return None


def _resolve_link(doctype: str, value: str, company: str | None = None) -> str | None:
	raw = (value or "").strip()
	if not raw:
		return None
	if frappe.db.exists(doctype, raw):
		return raw
	meta = frappe.get_meta(doctype)
	title_field = meta.get("title_field")
	if title_field:
		filters = {title_field: raw}
		if company and meta.has_field("company"):
			filters["company"] = company
		found = frappe.db.get_value(doctype, filters, "name")
		if found:
			return found
		found = frappe.db.get_value(doctype, {title_field: raw}, "name")
		if found:
			return found
	if doctype == "Department" and company:
		abbr = frappe.db.get_value("Company", company, "abbr")
		if abbr and frappe.db.exists("Department", f"{raw} - {abbr}"):
			return f"{raw} - {abbr}"
	if doctype == "Customer":
		return frappe.db.get_value("Customer", {"customer_name": raw}, "name")
	if doctype == "Employee Grade" and frappe.db.exists("Employee Grade", raw):
		return raw
	return None


def _resolve_employee(value: str, company: str | None = None) -> str | None:
	raw = (value or "").strip()
	if not raw:
		return None
	if frappe.db.exists("Employee", raw):
		return raw
	filters = {"employee_name": raw}
	if company:
		filters["company"] = company
		found = frappe.db.get_value("Employee", filters, "name")
		if found:
			return found
	found = frappe.db.get_value("Employee", {"employee_name": raw}, "name")
	if found:
		return found
	for fieldname in ("company_email", "personal_email", "employee_number"):
		found = frappe.db.get_value("Employee", {fieldname: raw}, "name")
		if found:
			return found
	return None


def _full_name(row: list[str], mapping: dict[int, str]) -> str:
	values = {fieldname: row[idx] if idx < len(row) else "" for idx, fieldname in mapping.items()}
	return _full_name_from_values(values)


def _full_name_from_values(values: dict) -> str:
	return " ".join(part for part in ((values.get("first_name") or "").strip(), (values.get("last_name") or "").strip()) if part)


def _load_bytes(filename: str | None, filedata: str | None, file_url: str | None) -> tuple[str, bytes]:
	if file_url:
		file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
		if not file_name:
			frappe.throw(_("Uploaded file was not found."))
		file_doc = frappe.get_doc("File", file_name)
		content = file_doc.get_content()
		if isinstance(content, str):
			content = content.encode("utf-8")
		return file_doc.file_name or filename or "import.xlsx", content
	if not filedata:
		frappe.throw(_("Upload a CSV or Excel file."))
	raw = filedata
	if isinstance(raw, str):
		prefix = raw[:80]
		if "," in prefix and "base64" in prefix.lower():
			raw = raw.split(",", 1)[1]
		try:
			raw = base64.b64decode(raw)
		except Exception:
			raw = raw.encode("utf-8")
	return filename or "import.csv", raw


def _read_csv(content: bytes) -> list[list]:
	text = content.decode("utf-8-sig", errors="replace")
	sample = text[:4096]
	try:
		dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
	except csv.Error:
		dialect = csv.excel
	reader = csv.reader(io.StringIO(text), dialect)
	return [list(row) for row in reader]


def _read_xlsx(content: bytes) -> list[list]:
	from openpyxl import load_workbook

	workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	sheet = workbook.active
	rows = []
	for row in sheet.iter_rows(values_only=True):
		rows.append(list(row))
	workbook.close()
	return rows


def _cell_str(value) -> str:
	if value is None:
		return ""
	if isinstance(value, datetime):
		return value.date().isoformat()
	if isinstance(value, date):
		return value.isoformat()
	if isinstance(value, bool):
		return "1" if value else "0"
	if isinstance(value, int):
		return str(value)
	if isinstance(value, float):
		if value.is_integer():
			return str(int(value))
		return format(value, "f").rstrip("0").rstrip(".")
	return str(value).strip()
